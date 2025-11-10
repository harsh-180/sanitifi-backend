# app/tasks.py
import os
import json
import pandas as pd
import subprocess
from concurrent.futures import ThreadPoolExecutor
from django.conf import settings
from celery import shared_task
from celery import current_task
import warnings
from openpyxl import load_workbook
from .models import User, Projects


@shared_task(bind=True)
def merge_chunks_task(self, upload_id):
    temp_root = os.path.join(settings.MEDIA_ROOT, 'temp', upload_id)
    meta_path = os.path.join(temp_root, 'meta.json')

    if not os.path.exists(meta_path):
        return {'error': 'meta.json not found'}

    with open(meta_path, 'r', encoding='utf-8') as f:
        meta = json.load(f)

    file_name = meta['fileName']
    total_chunks = int(meta['totalChunks'])
    user_id = meta['user_id']
    project_name = meta.get('project_name')
    project_id = meta.get('project_id')
    file_type = meta.get('file_type')

    merged_path = os.path.join(temp_root, file_name)

    # Progress Setup 
    def update_progress(current, total):
        percent = int((current / total) * 100)
        self.update_state(state='PROGRESS', meta={'percent': percent})

    # Merge chunks using ThreadPoolExecutor 
    def read_chunk(i):
        part_path = os.path.join(temp_root, f"{i}.part")
        with open(part_path, 'rb') as p:
            return p.read()

    with open(merged_path, 'wb') as out:
        with ThreadPoolExecutor(max_workers=1) as executor:
            for i, chunk_data in enumerate(executor.map(read_chunk, range(total_chunks))):
                out.write(chunk_data)
                update_progress(i + 1, total_chunks)

    # Move merged file into final directory 
    user = User.objects.get(id=user_id)
    if project_id:
        project = Projects.objects.get(id=project_id)
    else:
        project, _ = Projects.objects.get_or_create(
            user=user, name=project_name, defaults={'kpi_file': [], 'media_file': []}
        )

    project_folder = os.path.join(settings.MEDIA_ROOT, f"user_{user.id}/project_{project.id}")
    os.makedirs(project_folder, exist_ok=True)

    base_subdir = 'kpi' if file_type == 'kpi' else 'media'
    file_basename = os.path.splitext(file_name)[0]
    file_dir = os.path.join(project_folder, base_subdir, file_basename)
    os.makedirs(file_dir, exist_ok=True)
    final_path = os.path.join(file_dir, file_name)
    os.replace(merged_path, final_path)

    # Optional Pandas cleanup 
    file_size_mb = os.path.getsize(final_path) / (1024 * 1024)
    file_ext = os.path.splitext(file_name)[1].lower()

    try:
        if file_ext == '.csv':
            # Try multiple encodings to handle different file encodings
            try:
                df = pd.read_csv(final_path, encoding='utf-8')
            except UnicodeDecodeError:
                try:
                    df = pd.read_csv(final_path, encoding='latin1')
                except UnicodeDecodeError:
                    df = pd.read_csv(final_path, encoding='cp1252')
            
            for col in df.columns:
                df[col] = pd.to_numeric(df[col], errors='ignore')
            df.to_csv(final_path, index=False, encoding='utf-8')
            
        elif file_ext in ['.xlsx', '.xls']:
            
            warnings.filterwarnings("ignore", category=UserWarning)
            print(f"Processing Excel file: {final_path}")

            try:
                # Get all sheet names quickly first
                wb = load_workbook(final_path, read_only=True, data_only=True)
                sheet_names = wb.sheetnames
                wb.close()
                print(f"Found sheets: {sheet_names}")
            except Exception as e:
                print(f" Could not read Excel sheets: {e}")
                raise

            # Read each sheet separately using pandas
            for sheet_name in sheet_names:
                try:
                    print(f" Reading sheet: {sheet_name}")
                    df = pd.read_excel(final_path, sheet_name=sheet_name, dtype=str, engine='openpyxl')

                    # Convert numeric values
                    for col in df.columns:
                        try:
                            df[col] = pd.to_numeric(df[col], errors='ignore')
                            if df[col].dtype == 'float64':
                                df[col] = df[col].round(2)
                        except Exception:
                            pass

                    # Save to CSV
                    sheet_path = os.path.join(file_dir, f"{sheet_name}.csv")
                    df.to_csv(sheet_path, index=False)
                    print(f"Converted sheet '{sheet_name}' → {sheet_path}")
                except Exception as e:
                    print(f" Skipping sheet {sheet_name}: {e}")
                    continue

            try:
                os.remove(final_path)
            except Exception as e:
                print(f" Could not remove Excel file: {e}")   
            
    except Exception as e:
        self.update_state(state='FAILURE', meta={'error': str(e)})
        print(f"Error while processing file {file_name}: {e}")   
        raise

    # Update project's file list
    if file_type == 'kpi':
        kpi_list = project.kpi_file if isinstance(project.kpi_file, list) else []
        if file_basename not in kpi_list:
            kpi_list.append(file_basename)
            project.kpi_file = kpi_list
    else:  # media
        media_list = project.media_file if isinstance(project.media_file, list) else []
        if file_basename not in media_list:
            media_list.append(file_basename)
            project.media_file = media_list

    project.save()

    # Commit to Git 
    try:
        if not os.path.exists(os.path.join(project_folder, ".git")):
            subprocess.run(["git", "init"], cwd=project_folder, check=False)
        subprocess.run(["git", "add", "-A"], cwd=project_folder, check=False)
        subprocess.run(["git", "commit", "-m", f"Merged {file_name}"], cwd=project_folder, check=False)
    except Exception:
        pass

    # Cleanup temporary folder
    for entry in os.listdir(temp_root):
        os.remove(os.path.join(temp_root, entry))
    os.rmdir(temp_root)

    return {'success': True, 'project_id': project.id, 'percent': 100}
